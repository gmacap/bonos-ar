#!/usr/bin/env python3
"""Extrae el ultimo relevamiento del REM del BCRA a rem.json.

El BCRA publica el historico completo del Relevamiento de Expectativas de
Mercado en una planilla sin fecha en el nombre. Pero el nombre cambio: hasta
agosto de 2026 era historico-relevamiento-expectativas-mercado.xlsx y desde el
relevamiento de septiembre la pagina del REM enlaza
relevamiento-expectativas-mercado-historico.xlsx. El viejo quedo publicado y
congelado en el 31/08, y como el script lo seguia leyendo, el REM de la app se
quedo un mes atras sin que nada fallara. Por eso ahora se busca el link en la
pagina, se prueban los dos nombres conocidos y gana el relevamiento mas nuevo.
De ahi salen las tres series que usa la app:

  - inflacion mensual (var. % mensual)
  - TAMAR de bancos privados (TNA; %)
  - tipo de cambio nominal mayorista ($/USD, en niveles, fin de mes)

Mas las anclas anuales a diciembre, que es lo unico que publica el REM mas
alla de los seis meses del tramo mensual.

Uso:
    python .github/scripts/rem.py              # escribe rem.json
    python .github/scripts/rem.py --dry-run    # lo imprime y no toca nada
"""

import argparse
import io
import json
import os
import re
import sys
import urllib.request
from datetime import datetime, timezone

BASE = "https://www.bcra.gob.ar"
PAGINA = BASE + "/PublicacionesEstadisticas/Relevamiento_Expectativas_de_Mercado.asp"
CARPETA = BASE + "/archivos/Pdfs/PublicacionesEstadisticas/informes/"
CONOCIDAS = [
    CARPETA + "relevamiento-expectativas-mercado-historico.xlsx",
    CARPETA + "historico-relevamiento-expectativas-mercado.xlsx",
]

HOJA = "Base de Datos Completa"

# (clave de salida, texto que identifica la Variable, Referencia del tramo mensual)
SERIES = [
    ("ipc",   "IPC nivel general",   "var. % mensual"),
    ("tamar", "TAMAR",               "TNA; %"),
    ("tcn",   "Tipo de cambio nominal", "$/USD"),
]

# El ancla anual llega como "var. % i.a.; dic-27" o "$/USD; dic-27": nos
# interesan solo las de diciembre, que cierran el ano calendario.
MESES_ES = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
            "jul": 7, "ago": 8, "sep": 9, "oct": 10, "nov": 11, "dic": 12}


def descargar(url, timeout=180):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def candidatas():
    """El link de la planilla historica que enlaza hoy la pagina del REM, mas
    los nombres conocidos, sin repetir."""
    urls = []
    try:
        html = descargar(PAGINA, timeout=60).decode("utf-8", "replace")
        for href in re.findall(r'href="([^"]+\.xlsx)"', html):
            if "historico" in href.lower() and "expectativas" in href.lower():
                urls.append(href if href.startswith("http") else BASE + href)
    except Exception as e:
        print(f"   (no se pudo leer la pagina del REM: {e})")
    for u in CONOCIDAS:
        if u not in urls:
            urls.append(u)
    return urls


def norm(s):
    """Minusculas y sin acentos, para comparar sin pelearse con el encoding."""
    s = str(s or "").lower()
    for a, b in (("a", "á"), ("e", "é"), ("i", "í"),
                 ("o", "ó"), ("u", "ú")):
        s = s.replace(b, a)
    return s


def mes_de(valor):
    """'YYYY-MM' a partir de una celda de Periodo (datetime o texto ISO)."""
    if isinstance(valor, datetime):
        return f"{valor.year:04d}-{valor.month:02d}"
    txt = str(valor or "")
    if len(txt) >= 7 and txt[4] == "-":
        return txt[:7]
    return None


def ancla_de(referencia):
    """'2027-12' si la referencia termina en un 'dic-27'; si no, None."""
    txt = norm(referencia)
    if ";" not in txt:
        return None
    cola = txt.rsplit(";", 1)[1].strip()
    if "-" not in cola:
        return None
    mes, anio = cola.split("-", 1)
    mes, anio = mes.strip()[:3], anio.strip()
    if mes not in MESES_ES or not anio.isdigit():
        return None
    if MESES_ES[mes] != 12:          # solo las de cierre de ano
        return None
    return f"20{anio[-2:]}-12"


def extraer(xlsx_bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    if HOJA not in wb.sheetnames:
        raise SystemExit(f"La hoja '{HOJA}' no esta en el archivo. Hojas: {wb.sheetnames}")
    ws = wb[HOJA]

    filas = []
    for i, fila in enumerate(ws.iter_rows(values_only=True)):
        if i < 2 or fila[0] is None:
            continue
        filas.append(fila)
    if not filas:
        raise SystemExit("La hoja vino vacia.")

    ultimo = max(mes_de(f[0]) or "" for f in filas)
    fechas = sorted({str(f[0])[:10] for f in filas if (mes_de(f[0]) or "") == ultimo})
    relevamiento = fechas[-1]

    out = {"relevamiento": relevamiento,
           "actualizado": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
           "anclas": {}}

    for clave, marca, ref_mensual in SERIES:
        marca_n, ref_n = norm(marca), norm(ref_mensual)
        mensuales, anclas = [], {}
        for f in filas:
            if str(f[0])[:10] != relevamiento:
                continue
            var, ref, per, mediana = norm(f[1]), str(f[2] or ""), f[3], f[4]
            if marca_n not in var or mediana is None:
                continue
            # El IPC nucleo comparte la marca del general: hay que descartarlo.
            if clave == "ipc" and "nucleo" in var:
                continue
            if norm(ref) == ref_n:
                mes = mes_de(per)
                if mes:
                    mensuales.append({"mes": mes, "v": round(float(mediana), 4)})
            else:
                a = ancla_de(ref)
                if a:
                    anclas[a] = round(float(mediana), 4)
        mensuales.sort(key=lambda x: x["mes"])
        if not mensuales:
            raise SystemExit(f"No salio ningun punto mensual para '{clave}'. "
                             f"Cambio el formato del archivo del BCRA?")
        out[clave] = mensuales
        out["anclas"][clave] = dict(sorted(anclas.items()))

    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true",
                    help="imprime el resultado y no escribe rem.json")
    ap.add_argument("--salida", default=None, help="ruta del json (default: rem.json en la raiz)")
    args = ap.parse_args()

    raiz = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    destino = args.salida or os.path.join(raiz, "rem.json")

    # Cada planilla candidata se baja y se parsea; gana el relevamiento mas
    # nuevo. Una que no baja o no tiene el formato esperado se saltea.
    datos = None
    for url in candidatas():
        print(f"-> Bajando {url}")
        try:
            d = extraer(descargar(url))
        except (Exception, SystemExit) as e:
            print(f"   salteada: {e}")
            continue
        print(f"   relevamiento {d['relevamiento']}")
        if datos is None or d["relevamiento"] > datos["relevamiento"]:
            datos = d
    if datos is None:
        raise SystemExit("Ninguna planilla del REM se pudo leer.")

    print(f"\n   Usado: relevamiento {datos['relevamiento']}")
    for clave, _, _ in SERIES:
        ms = datos[clave]
        anclas = datos["anclas"][clave]
        print(f"   {clave:<6} {len(ms)} meses  {ms[0]['mes']} -> {ms[-1]['mes']}"
              f"   primero={ms[0]['v']}  anclas={anclas or '-'}")

    texto = json.dumps(datos, ensure_ascii=False, indent=2) + "\n"

    if args.dry_run:
        print("\n" + texto)
        return

    # Si solo cambio el sello de tiempo, no reescribimos: el workflow commitea
    # cuando el archivo cambia, y un commit diario que solo mueve la hora seria
    # ruido en el historial.
    if os.path.exists(destino):
        try:
            with open(destino, encoding="utf-8") as fh:
                previo = json.load(fh)
            sin_sello = {k: v for k, v in datos.items() if k != "actualizado"}
            if {k: v for k, v in previo.items() if k != "actualizado"} == sin_sello:
                print(f"\nSin cambios respecto de {destino}. No se reescribe.")
                return
        except Exception:
            pass

    with open(destino, "w", encoding="utf-8", newline="\r\n") as fh:
        fh.write(texto)
    print(f"\nEscrito {destino}")


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)
